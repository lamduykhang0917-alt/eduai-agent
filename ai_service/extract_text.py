import os

import pdfplumber
from pptx import Presentation

SUBJECT_MARGINS = {
    "ADO.NET": {"top": 0.0, "bottom": 0.0},
    "CAU_TRUC_DU_LIEU": {"top": 0.08, "bottom": 0.08},
    "CO_SO_DU_LIEU": {"top": 0.08, "bottom": 0.08},
    "DIEN_TOAN_DAM_MAY": {"top": 0.08, "bottom": 0.08},
    "DO_HOA_MAY_TINH": {"top": 0.08, "bottom": 0.10},
    "GIAI_THUAT": {"top": 0.08, "bottom": 0.08},
    "HE_QUAN_TRI_CO_SO_DU_LIEU": {"top": 0.08, "bottom": 0.08},
    "KIEN_TRUC_MAY_TINH": {"top": 0.10, "bottom": 0.08},
    "KY_THUAT_LAP_TRINH": {"top": 0.10, "bottom": 0.08},
    "MANG_MAY_TINH": {"top": 0.0, "bottom": 0.0},
    "MAY_HOC": {"top": 0.0, "bottom": 0.0},
    "NHAP_MON_KHMT": {"top": 0.0, "bottom": 0.00},
    "LAP_TRINH_HUONG_DOI_TUONG": {"top": 0.08, "bottom": 0.08},
    "TRI_TUE_NHAN_TAO": {"top": 0.08, "bottom": 0.08},
    "TUONG_TAC_NGUOI_MAY": {"top": 0.00, "bottom": 0.0},
    "THU_THAP_TIEN_XU_LI_DU_LIEU": {"top": 0.08, "bottom": 0.08},
    "XU_LI_ANH": {"top": 0.0, "bottom": 0.0},
    "XU_LI_NGON_NGU_TU_NHIEN": {"top": 0.08, "bottom": 0.08},
}

DEFAULT_MARGIN = {"top": 0.05, "bottom": 0.05}

PAGE_BREAK = "\n\n---PAGE_BREAK---\n\n"


def extract_pdf_pages(pdf_path: str, top_margin: float, bottom_margin: float) -> list:
    pages = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            width, height = page.width, page.height
            bbox = (0, height * top_margin, width, height * (1 - bottom_margin))
            cropped_page = page.within_bbox(bbox)
            text = cropped_page.extract_text()
            pages.append(text or "")
    return pages


def extract_pptx_pages(pptx_path: str, top_margin: float, bottom_margin: float) -> list:
    prs = Presentation(pptx_path)
    slide_height = prs.slide_height
    pages = []
    for slide in prs.slides:
        slide_text = []
        shapes = sorted(
            slide.shapes,
            key=lambda s: (s.top if hasattr(s, "top") and s.top is not None else 0,
                            s.left if hasattr(s, "left") and s.left is not None else 0),
        )
        for shape in shapes:
            if not shape.has_text_frame:
                continue
            if hasattr(shape, "top") and shape.top is not None and slide_height > 0:
                if shape.top < slide_height * top_margin:
                    continue
                if shape.top > slide_height * (1 - bottom_margin):
                    continue
            text = shape.text.strip()
            if text:
                slide_text.append(text)
        pages.append("\n".join(slide_text))
    return pages


def extract_file_pages(file_path: str, subject: str) -> list:
    margin = SUBJECT_MARGINS.get(subject, DEFAULT_MARGIN)
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_pdf_pages(file_path, margin["top"], margin["bottom"])
    if ext == ".pptx":
        return extract_pptx_pages(file_path, margin["top"], margin["bottom"])
    raise ValueError(f"Không hỗ trợ định dạng: {ext}")
