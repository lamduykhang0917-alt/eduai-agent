"""Đọc nội dung chữ từ tệp tài liệu (PDF, DOCX, TXT) để đưa vào kho tri thức của chatbot."""

import io
import re

MAX_BYTES = 15 * 1024 * 1024
ALLOWED = {"pdf", "docx", "txt"}
CHUNK_CHARS = 1500


class ExtractError(Exception):
    pass


def file_type_of(filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED:
        raise ExtractError("Chỉ hỗ trợ tệp PDF, DOCX hoặc TXT")
    return ext


def _split_long(text: str, limit: int = CHUNK_CHARS):
    """Cắt đoạn dài thành các khúc ≤ limit ký tự, ưu tiên cắt ở hết đoạn/câu."""
    text = text.strip()
    if len(text) <= limit:
        return [text] if text else []
    parts, current = [], ""
    for para in re.split(r"\n{1,}", text):
        if len(current) + len(para) + 1 > limit and current:
            parts.append(current.strip())
            current = ""
        while len(para) > limit:
            cut = para.rfind(". ", 0, limit)
            cut = cut + 1 if cut > limit // 2 else limit
            parts.append(para[:cut].strip())
            para = para[cut:]
        current += para + "\n"
    if current.strip():
        parts.append(current.strip())
    return parts


def _pdf_pages(data: bytes):
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        return [(i, page.extract_text() or "") for i, page in enumerate(reader.pages, 1)]
    except Exception as exc:  # tệp hỏng / có mật khẩu
        raise ExtractError("Không đọc được tệp PDF (tệp hỏng hoặc có mật khẩu)") from exc


def _docx_pages(data: bytes):
    import docx
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractError("Không đọc được tệp DOCX") from exc
    lines = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return _group_lines(lines)


def _txt_pages(data: bytes):
    for enc in ("utf-8-sig", "utf-16", "cp1258", "cp1252"):
        try:
            text = data.decode(enc)
            break
        except UnicodeError:
            continue
    else:
        raise ExtractError("Không đọc được bảng mã của tệp TXT")
    return _group_lines(text.splitlines())


def _group_lines(lines):
    """Gom các dòng thành từng 'trang' khoảng 3000 ký tự để hiển thị/tìm kiếm theo phần."""
    pages, current, size = [], [], 0
    for line in lines:
        current.append(line)
        size += len(line) + 1
        if size >= 3000:
            pages.append("\n".join(current))
            current, size = [], 0
    if current:
        pages.append("\n".join(current))
    return list(enumerate(pages, 1))


def extract_chunks(filename: str, data: bytes):
    """Trả về (file_type, [(page_number, text), ...]) với mỗi phần tử là một đoạn ≤ 1500 ký tự."""
    if len(data) > MAX_BYTES:
        raise ExtractError("Tệp quá lớn (tối đa 15 MB)")
    ftype = file_type_of(filename)
    pages = {"pdf": _pdf_pages, "docx": _docx_pages, "txt": _txt_pages}[ftype](data)
    chunks = []
    for page_no, text in pages:
        for piece in _split_long(text):
            chunks.append((page_no, piece))
    if not chunks:
        raise ExtractError("Không tìm thấy chữ trong tệp (có thể là bản scan hoặc tệp rỗng)")
    return ftype, chunks
